"""Dosya ayrıştırma ve kaynak konumunu koruyan karakter tabanlı parçalama."""
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from pypdf import PdfReader
from docx import Document as WordDocument


class DocumentError(ValueError):
    pass


@dataclass
class Section:
    location: str
    text: str


def clean_text(text):
    text = text.replace("\x00", "").replace("\r\n", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_document(path: Path, filename: str, max_pages=300):
    suffix = Path(filename).suffix.lower()
    sections, warnings = [], []
    if suffix == ".pdf":
        with path.open("rb") as file:
            if file.read(5) != b"%PDF-":
                raise DocumentError("Dosya geçerli bir PDF değil.")
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentError("Şifreli PDF desteklenmiyor; şifreyi kaldırıp yeniden yükle.")
        if len(reader.pages) > max_pages:
            raise DocumentError(f"En fazla {max_pages} sayfa yüklenebilir; dosyayı böl.")
        empty_pages = []
        for number, page in enumerate(reader.pages, 1):
            text = clean_text(page.extract_text() or "")
            if len(text) < 20:
                empty_pages.append(number)
            if text:
                sections.append(Section(f"PDF sayfa {number}", text))
        if empty_pages:
            warnings.append("Metni az/olmayan sayfalar (OCR gerekebilir): " + ", ".join(map(str, empty_pages[:40])))
        if not sections or sum(len(s.text) for s in sections) < 40:
            raise DocumentError("Metin çıkarılamadı. Taranmış PDF için önce yerel OCR uygula (rehbere bak).")
    elif suffix == ".docx":
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(x.file_size for x in entries) > 50 * 1024 * 1024:
                raise DocumentError("DOCX açılmış boyutu güvenli sınırı aşıyor.")
        document = WordDocument(path)
        for number, paragraph in enumerate(document.paragraphs, 1):
            if text := clean_text(paragraph.text):
                sections.append(Section(f"DOCX paragraf {number}", text))
        for table_index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                if text := clean_text(" | ".join(cell.text for cell in row.cells)):
                    sections.append(Section(f"Tablo {table_index}, satır {row_index}", text))
        warnings.append("DOCX konumları paragraf/tablo referansıdır; sayfa numarası değildir. Dipnot, resim ve metin kutuları alınmaz.")
    elif suffix in {".txt", ".md"}:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentError("TXT/MD UTF-8 kodlamasında olmalı.") from exc
        offset = 1
        for paragraph in re.split(r"\n\s*\n", text):
            if clean_text(paragraph):
                sections.append(Section(f"Metin blok {offset}", clean_text(paragraph)))
                offset += 1
    else:
        raise DocumentError("Yalnızca PDF, DOCX, TXT ve MD desteklenir.")
    if not sections:
        raise DocumentError("Dosyada işlenebilir metin yok.")
    if sum(len(x.text) for x in sections) > 1_500_000:
        raise DocumentError("Çıkarılan metin çok büyük; dosyayı böl.")
    return sections, " ".join(warnings)[:500]


def chunk_sections(sections, size=1600, overlap=220):
    if size < 1 or not 0 <= overlap < size:
        raise ValueError("Geçersiz parça/örtüşme ayarı.")
    chunks = []
    for section in sections:
        start = 0
        while start < len(section.text):
            end = min(start + size, len(section.text))
            if end < len(section.text):
                cut = section.text.rfind(" ", start + size // 2, end)
                if cut > start:
                    end = cut
            text = section.text[start:end].strip()
            if text:
                chunks.append(Section(section.location, text))
            if end == len(section.text):
                break
            start = max(start + 1, end - overlap)
    return chunks

