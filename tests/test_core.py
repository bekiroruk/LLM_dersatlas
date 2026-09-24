import tempfile
import unittest
import os
from pathlib import Path
from types import SimpleNamespace
from app.config import Settings, PROJECT_ROOT
from app.ranking import tokenize, bm25, reciprocal_rank_fusion
from app.passwords import hash_password, check_password, token_hash
from app.citations import valid_citations
from app.ingestion import Section, DocumentError, chunk_sections, extract_document, clean_text
from app.rag import (
    _declared_source_ids,
    _focused_vegetation_evidence,
    _focused_climate_aspect_evidence,
    _normalize_citation_shapes,
    _retrieval_queries,
    _evidence_coverage_keys,
    _seasonal_climate_phrase,
    _precipitation_relation_quality,
    _is_question_catalog,
    _comparison_search_queries,
    _sources_are_relevant,
    _generic_comparison_evidence,
    _answer_has_source_support,
    _range_claims_supported,
    _duration_claims_supported,
    _agency_claims_supported,
    _answer_claim_units,
    _contains_source_meta_claim,
)
from pypdf import PdfWriter
from docx import Document as WordDocument


class CoreTests(unittest.TestCase):
    def test_question_catalog_is_not_declarative_evidence(self):
        catalog = (
            "Türkiye’de yasama yetkisi hangi organa aittir? "
            "Türkiye’de yürütme yetkisi ve görevi kime aittir? "
            "Yargı yetkisi kimlerce kullanılır?"
        )
        answered = (
            "Soru: Yasama yetkisi kime aittir? "
            "Cevap: Yasama yetkisi TBMM'ye aittir."
        )
        self.assertTrue(_is_question_catalog(catalog))
        self.assertFalse(_is_question_catalog(answered))

    def test_runtime_data_paths_do_not_depend_on_terminal_directory(self):
        previous = Path.cwd()
        with tempfile.TemporaryDirectory() as directory:
            try:
                os.chdir(directory)
                settings = Settings(
                    _env_file=None,
                    data_dir=Path("data-test-root"),
                    database_url="sqlite:///./data-test-root/dersatlas.db",
                )
            finally:
                os.chdir(previous)

        expected = (PROJECT_ROOT / "data-test-root").resolve()
        self.assertEqual(settings.data_dir, expected)
        self.assertEqual(
            settings.database_url,
            "sqlite:///" + (expected / "dersatlas.db").as_posix(),
        )

    def test_comparison_retrieval_splits_both_sides(self):
        question = "Karadeniz ve Akdeniz iklimi açısından bitki örtüsü nasıl farklı?"
        self.assertEqual(
            _retrieval_queries(question),
            [
                question,
                "Karadeniz iklimi bitki örtüsü nasıl farklı",
                "Akdeniz iklimi bitki örtüsü nasıl farklı",
                "Karadeniz iklimi doğal bitki örtüsü flora bitki varlığı baskın görünüm",
                "Akdeniz iklimi doğal bitki örtüsü flora bitki varlığı baskın görünüm",
            ],
        )

    def test_generic_shared_head_comparison_splits_both_sides(self):
        question = "Tanzimat ve Islahat fermanlarının farkları nelerdir?"
        self.assertEqual(
            _comparison_search_queries(question),
            [
                "Tanzimat fermanlarının farkları nelerdir",
                "Islahat fermanlarının farkları nelerdir",
            ],
        )
        self.assertEqual(
            _retrieval_queries(question),
            [
                question,
                "Tanzimat fermanlarının farkları nelerdir",
                "Islahat fermanlarının farkları nelerdir",
            ],
        )

    def test_generic_comparison_accepts_separate_evidence_for_each_side(self):
        question = "Tanzimat ve Islahat fermanlarının farkları nelerdir?"
        tanzimat = {
            "chunk_id": "tanzimat",
            "text": (
                "Tanzimat Fermanı 3 Kasım 1839 tarihinde ilan edildi. "
                "Can, mal ve namus güvenliğini düzenledi."
            ),
        }
        islahat = {
            "chunk_id": "islahat",
            "text": (
                "Islahat Fermanı 1856 yılında ilan edildi. "
                "Gayrimüslim tebaanın haklarını genişletti."
            ),
        }
        self.assertTrue(_sources_are_relevant(question, [tanzimat, islahat]))
        self.assertFalse(_sources_are_relevant(question, [tanzimat]))

    def test_generic_comparison_prefers_exclusive_pages_over_shared_summary(self):
        question = "Tanzimat ve Islahat fermanlarının farkları nelerdir?"
        shared = {
            "chunk_id": "shared",
            "text": (
                "Tanzimat Fermanı'nda Mustafa Reşit Paşa etkilidir. "
                "Islahat Fermanı'nda Âli ve Fuat Paşalar etkilidir."
            ),
        }
        tanzimat = {
            "chunk_id": "tanzimat",
            "text": (
                "Tanzimat Fermanı 3 Kasım 1839 tarihinde ilan edildi. "
                "Can, mal ve namus güvenliğini düzenledi."
            ),
        }
        islahat = {
            "chunk_id": "islahat",
            "text": (
                "Islahat Fermanı 1856 yılında ilan edildi. "
                "Gayrimüslim tebaanın haklarını genişletti."
            ),
        }
        sides = _generic_comparison_evidence(
            question,
            [shared, tanzimat, islahat],
            per_side=1,
        )
        self.assertEqual(
            [side[0][0]["chunk_id"] for side in sides],
            ["tanzimat", "islahat"],
        )

    def test_generic_comparison_combines_claims_only_from_selected_side_sources(self):
        question = "Tanzimat ve Islahat fermanlarının farkları nelerdir?"
        sources = [
            {
                "text": (
                    "Tanzimat Fermanı 3 Kasım 1839 tarihinde ilan edildi. "
                    "Can, mal ve namus güvenliğini düzenledi."
                ),
            },
            {
                "text": (
                    "Islahat Fermanı 1856 yılında ilan edildi. "
                    "Gayrimüslim tebaanın haklarını genişletti."
                ),
            },
        ]
        grounded = (
            "Tanzimat Fermanı 1839 yılında ilan edilip can, mal ve namus "
            "güvenliğini düzenlerken Islahat Fermanı 1856 yılında ilan "
            "edilmiş ve gayrimüslim tebaanın haklarını genişletmiştir."
        )
        self.assertTrue(
            _answer_has_source_support(grounded, question, sources)
        )
        self.assertFalse(
            _answer_has_source_support(
                grounded.replace("1856", "1908"),
                question,
                sources,
            )
        )

    def test_period_range_cannot_become_edict_announcement_range(self):
        source = (
            "Tanzimat Dönemi 1839-1876 arasındadır. "
            "Tanzimat Fermanı'nda Mustafa Reşit Paşa etkilidir."
        )
        self.assertFalse(
            _range_claims_supported(
                "Tanzimat Fermanı 1839-1876 yılları arasında ilan edilmiştir.",
                [source],
            )
        )
        self.assertTrue(
            _range_claims_supported(
                "Tanzimat Dönemi 1839-1876 arasındadır.",
                [source],
            )
        )

    def test_period_end_cannot_become_edict_duration(self):
        period = (
            "Tanzimat Dönemi 1839 Tanzimat Fermanı'nın ilanından "
            "1876 I. Meşrutiyet'in ilanına kadar geçen süreçtir."
        )
        claim = (
            "Tanzimat Fermanı 1839 yılında ilan edilmiştir ve 1876 "
            "yılına kadar devam etmiştir."
        )
        self.assertFalse(_duration_claims_supported(claim, [period]))
        self.assertTrue(_duration_claims_supported(
            claim,
            ["Tanzimat Fermanı 1876 yılına kadar devam etmiştir."],
        ))

    def test_influential_person_cannot_become_announcing_agent(self):
        claim = "Tanzimat Fermanı Mustafa Reşit Paşa tarafından ilan edildi."
        self.assertFalse(_agency_claims_supported(
            claim,
            [
                "Tanzimat Fermanı'nda Mustafa Reşit Paşa etkilidir. "
                "Tanzimat Fermanı 1839 yılında ilan edildi."
            ],
        ))
        self.assertTrue(_agency_claims_supported(
            claim,
            ["Tanzimat Fermanı Mustafa Reşit Paşa tarafından ilan edildi."],
        ))

    def test_conjoined_independent_claims_are_checked_separately(self):
        units = _answer_claim_units(
            "Islahat Fermanı ile gayrimüslimlere geniş haklar verilmiştir "
            "ve Müslümanlarla gayrimüslimler arasındaki eşitlik artırılmıştır."
        )
        self.assertEqual(len(units), 2)
        self.assertIn("geniş haklar", units[0])
        self.assertIn("eşitlik", units[1])

    def test_source_heading_and_meta_language_are_not_answer_claims(self):
        self.assertTrue(
            _contains_source_meta_claim(
                "İlgili kaynaklarda '118. CİZYE VERGİSİNİN "
                "KALDIRILMASIYLA İLİŞKİ' gibi detaylar bulunmaktadır."
            )
        )
        self.assertFalse(
            _contains_source_meta_claim(
                "Islahat Fermanı gayrimüslimlere yeni haklar tanımıştır."
            )
        )

    def test_mixed_comparison_adds_dedicated_rainfall_searches(self):
        question = (
            "Karadeniz ve Akdeniz iklimlerini yağış rejimleri ve doğal "
            "bitki örtüleri bakımından karşılaştır."
        )

    def test_required_evidence_is_found_outside_similarity_shortlist(self):
        question = (
            "Karadeniz ve Akdeniz iklimlerini yağış rejimleri ve doğal "
            "bitki örtüleri bakımından karşılaştır."
        )
        allowed = {
            f"distractor-{index:03d}": (
                SimpleNamespace(
                    id=f"distractor-{index:03d}",
                    text=(
                        "Karadeniz ve Akdeniz iklim bölgeleri için yağış "
                        f"ve bitki örtüsü alıştırması {index}."
                    ),
                ),
                f"distractor-{index}.pdf",
                "Coğrafya",
            )
            for index in range(80)
        }
        allowed["vegetation-real"] = (
            SimpleNamespace(
                id="vegetation-real",
                text=(
                    "11. TÜRKİYE'NİN BİTKİ VARLIĞI\n"
                    "Flora bölgesi Türkiye’de yayılışı Baskın görünüm\n"
                    "Avrupa-Sibirya Marmara’nın kuzeyi ve Karadeniz kıyı "
                    "kuşağı Nemli ormanlar\nAkdeniz\nGüney Marmara, Ege, "
                    "Akdeniz ve Güneydoğu’nun batısına uzanan Akdeniz "
                    "iklim sahaları\nKızılçam, maki ve kuraklığa dayanıklı "
                    "Akdeniz türleri"
                ),
            ),
            "4 - TÜRKİYE SU, TOPRAK VE BİTKİ.pdf",
            "Coğrafya",
        )
        allowed["rainfall-real"] = (
            SimpleNamespace(
                id="rainfall-real",
                text=(
                    "Yayılış Gürcistan sınırından Bulgaristan sınırına kadar "
                    "Karadeniz kıyıları; Yaz Serin ve yağışlı Kış Ilık ve "
                    "yağışlı En fazla yağış Sonbahar\nGüney Marmara, Ege "
                    "kıyıları, Akdeniz kıyıları; Yaz Sıcak ve kurak Kış Ilık "
                    "ve yağışlı En fazla yağış Kış"
                ),
            ),
            "3 - TÜRKİYE İKLİMİ.pdf",
            "Coğrafya",
        )

        keys, covered, required = _evidence_coverage_keys(question, allowed)

        self.assertEqual(keys, ["vegetation-real", "rainfall-real"])
        self.assertEqual((covered, required), (4, 4))
        self.assertEqual(
            _retrieval_queries(question),
            [
                "Karadeniz iklimi yağış rejimi yağışların mevsimlere dağılışı en fazla yağış en az yağış",
                "Akdeniz iklimi yağış rejimi yağışların mevsimlere dağılışı en fazla yağış en az yağış",
                "Karadeniz iklimi doğal bitki örtüsü flora bitki varlığı baskın görünüm",
                "Akdeniz iklimi doğal bitki örtüsü flora bitki varlığı baskın görünüm",
            ],
        )

    def test_attached_ile_comparison_splits_both_climates(self):
        question = (
            "Akdeniz iklimiyle Karadeniz iklimini yağış düzeni ve "
            "bitki örtüsü yönünden kıyaslar mısın?"
        )
        queries = _retrieval_queries(question)
        self.assertEqual(len(queries), 4)
        self.assertTrue(any(query.startswith("Akdeniz iklimi yağış rejimi") for query in queries))
        self.assertTrue(any(query.startswith("Karadeniz iklimi yağış rejimi") for query in queries))
        self.assertIn(
            "Akdeniz iklimi doğal bitki örtüsü flora bitki varlığı baskın görünüm",
            queries,
        )
        self.assertIn(
            "Karadeniz iklimi doğal bitki örtüsü flora bitki varlığı baskın görünüm",
            queries,
        )

    def test_false_year_question_adds_a_year_neutral_verification_query(self):
        question = "Islahat Fermanı 1876 yılında mı ilan edildi?"
        self.assertEqual(
            _retrieval_queries(question),
            [
                question,
                "Islahat Fermanı yılında ilan edildi",
                "Islahat Fermanı ilan tarihi hangi yıl",
            ],
        )

    def test_rainfall_evidence_rejects_answer_key_and_wildfire_text(self):
        question = (
            "Karadeniz ve Akdeniz iklimlerini yağış rejimleri ve doğal "
            "bitki örtüleri bakımından karşılaştır."
        )
        distractors = (
            "1 E 9 III-IV 17 Doğu Karadeniz güney yamacı 18 Kahverengi orman 19 Nemlilik ve yağış",
            "Türkiye’de orman yangını hassasiyeti Akdeniz ikliminin görüldüğü kıyılarda yüksektir. Yaz sıcaklığı ve kuraklığı etkilidir.",
        )
        for text in distractors:
            with self.subTest(text=text):
                self.assertIsNone(
                    _focused_climate_aspect_evidence(question, text, "precipitation")
                )

    def test_flattened_climate_matrix_is_not_assigned_to_either_subject(self):
        question = (
            "Karadeniz ve Akdeniz iklimlerini yağış rejimleri ve doğal "
            "bitki örtüleri bakımından karşılaştır."
        )
        flattened = (
            "İklim tipi Karadeniz Akdeniz Sert karasal\n"
            "Yaz sıcak ve kurak\n"
            "Kış soğuk ve kar yağışlı\n"
            "Kış sıcaklığı 0 üstü yaklaşık 8-10 0 üstü yaklaşık 5 "
            "0 çevresi/altı Belirgin eksi\n"
            "Yaz kuraklığı Belirgin Yok Belirgin Yok\n"
            "Yağış rejimi Düzensiz Düzenli Düzensiz Düzensiz"
        )

        self.assertIsNone(
            _focused_climate_aspect_evidence(
                question,
                flattened,
                "precipitation",
            )
        )
        self.assertEqual(
            _seasonal_climate_phrase(
                "Akdeniz iklimi Kış sıcaklığı 0 üstü yaklaşık 8-10 "
                "0 üstü yaklaşık 5 0 çevresi/altı Belirgin eksi "
                "Yaz kuraklığı Belirgin Yok Belirgin Yok Yağış rejimi "
                "Düzensiz Düzenli Düzensiz Düzensiz",
                "akdeniz",
                "precipitation",
            ),
            "",
        )

    def test_nearby_soil_water_and_slope_rainfall_are_not_climate_regimes(self):
        question = (
            "Karadeniz ve Akdeniz iklimlerini yağış rejimleri ve doğal "
            "bitki örtüleri bakımından karşılaştır."
        )
        unrelated = (
            "Akdeniz Antalya, Mersin, İskenderun\n"
            "Fethiye Körfezi Ege-Akdeniz geçiş alanında yorumlanabilir.\n"
            "YERALTI SULARI VE KAYNAKLAR\n"
            "Vadi-yamaç kaynağı Yağıştan beslenir; rejimi düzensizdir.",
            "Kahverengi orman Orman örtüsü altında; özellikle\n"
            "Karadeniz ve diğer nemli kıyı ormanları\n"
            "Yağışla yıkanmış; tuz ve kireç az.",
            "Doğu Karadeniz kuzey yamacı Karadeniz'e dönüktür; nemli hava "
            "yükselip bol yağış bırakır. Güney yamaçta yağış azalır.",
        )

        for text in unrelated:
            with self.subTest(text=text):
                self.assertIsNone(
                    _focused_climate_aspect_evidence(
                        question,
                        text,
                        "precipitation",
                    )
                )

    def test_broad_comparison_requires_complete_rainfall_relation(self):
        question = (
            "Karadeniz ve Akdeniz iklimlerini yağış rejimleri ve doğal "
            "bitki örtüleri bakımından karşılaştır."
        )
        incomplete = (
            "Karadeniz iklimi\nYaz sıcak ve kurak\n"
            "Kış soğuk ve kar yağışlı.",
            "Akdeniz iklimi orta kuşak ve mutlak konumla ilişkilidir; "
            "en fazla yağış kışın cephelerle düşer.",
        )
        for text in incomplete:
            with self.subTest(text=text):
                self.assertIsNone(
                    _focused_climate_aspect_evidence(
                        question,
                        text,
                        "precipitation",
                    )
                )

        self.assertEqual(
            _precipitation_relation_quality(
                "Karadeniz kıyıları Yaz serin ve yağışlı Kış ılık ve "
                "yağışlı En fazla yağış sonbahar"
            )[0],
            1,
        )
        self.assertEqual(
            _precipitation_relation_quality(
                "Karadeniz ikliminde yağış yıl boyunca düzenlidir."
            )[0],
            1,
        )

    def test_rainfall_summary_drops_clipped_cells_and_exam_notes(self):
        karadeniz = (
            "Karadeniz iklimi yağış rejimi: ’de azdır. Tuzak / not: "
            "‘Yaz yağışlı’ ile ‘yaz kuraklığı yok’ aynı şey değildir; "
            "Karadeniz yıl boyu yağışlıdır."
        )
        akdeniz = (
            "Akdeniz iklimi yağış rejimi: kış yağışlı; yaz kurak, "
            "zeytin, rejim düzensiz."
        )

        self.assertEqual(
            _seasonal_climate_phrase(karadeniz, "karadeniz", "precipitation"),
            "yıl boyu yağışlıdır",
        )
        self.assertEqual(
            _seasonal_climate_phrase(akdeniz, "akdeniz", "precipitation"),
            "kış yağışlı; yaz kurak; rejim düzensizdir",
        )

    def test_vegetation_table_row_is_kept_as_one_evidence_relation(self):
        text = """11. TÜRKİYE'NİN BİTKİ VARLIĞI
11.1. Flora bölgeleri
Flora bölgesi Türkiye'de yayılışı Baskın görünüm
Avrupa-Sibirya Marmara'nın kuzeyi ve Karadeniz kıyı
kuşağı Nemli ormanlar
Akdeniz
Güney Marmara, Ege, Akdeniz ve
Güneydoğu'nun batısına uzanan
Akdeniz iklim sahaları
Kızılçam, maki ve kuraklığa dayanıklı Akdeniz türleri
Relikt Karadeniz'de kızılçam; Akdeniz'de kayın."""
        focus = _focused_vegetation_evidence(
            "Karadeniz ve Akdeniz iklimi açısından bitki örtüsü nasıl farklı?",
            text,
        )
        self.assertEqual(focus["covered"], ("karadeniz", "akdeniz"))
        self.assertIn("Karadeniz kıyı kuşağı Nemli ormanlar", focus["text"])
        self.assertIn("Akdeniz iklim sahaları Kızılçam, maki", focus["text"])
        self.assertNotIn("Relikt", focus["text"])

    def test_turkish_case(self):
        self.assertEqual(tokenize("ISLAHAT İSTANBUL ve 1839"), ["ıslahat", "istanbul", "1839"])

    def test_bm25_exact_date(self):
        result = bm25("1839 Tanzimat", [("a", "1856 Islahat Fermanı"), ("b", "1839 Tanzimat Fermanı")])
        self.assertEqual(result[0][0], "b")

    def test_no_lexical_match(self):
        self.assertEqual(bm25("kuantum", [("a", "Tanzimat Fermanı")]), [])

    def test_stop_words(self):
        self.assertEqual(bm25("ve veya için", [("a", "ve veya için")]), [])

    def test_rrf_consensus(self):
        result = reciprocal_rank_fusion([("a", .9), ("b", .8)], [("b", 10)])
        self.assertEqual(result[0][0], "b")
        self.assertLess(result[0][1], .1)

    def test_chunk_source_not_lost(self):
        parts = chunk_sections([Section("PDF sayfa 3", "Ankara " * 100), Section("PDF sayfa 4", "Meclis " * 100)], 100, 15)
        self.assertGreater(len(parts), 2)
        self.assertEqual({p.location for p in parts}, {"PDF sayfa 3", "PDF sayfa 4"})
        self.assertTrue(all(len(p.text) <= 100 for p in parts))

    def test_chunk_overlap_long_word(self):
        parts = chunk_sections([Section("blok 1", "a" * 301)], 100, 20)
        self.assertEqual(len(parts), 4)
        self.assertEqual(len(parts[-1].text), 61)

    def test_invalid_chunking(self):
        with self.assertRaises(ValueError):
            chunk_sections([], 100, 100)

    def test_text_cleanup(self):
        self.assertEqual(clean_text("A  B\r\n\n\nC\x00"), "A B\n\nC")

    def test_utf8_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "not.md"
            path.write_text("# Tarih\n\nİstanbul ve Millî Mücadele", encoding="utf-8")
            sections, _ = extract_document(path, path.name)
            self.assertEqual(len(sections), 2)
            self.assertIn("İstanbul", sections[1].text)

    def test_non_utf8_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "not.txt"
            path.write_bytes(b"\xff\xff")
            with self.assertRaises(DocumentError):
                extract_document(path, path.name)

    def test_blank_pdf_ocr_error(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "taranmis.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            writer.write(path)
            with self.assertRaisesRegex(DocumentError, "OCR"):
                extract_document(path, path.name)

    def test_docx_table_location(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "not.docx"
            document = WordDocument()
            document.add_paragraph("Tarih notu")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = "1839", "Tanzimat"
            document.save(path)
            sections, warning = extract_document(path, path.name)
            self.assertTrue(any(s.location == "Tablo 1, satır 1" and "1839" in s.text for s in sections))
            self.assertIn("sayfa numarası değildir", warning)

    def test_password_salt_and_verification(self):
        first = hash_password("Güvenli-parola-123")
        second = hash_password("Güvenli-parola-123")
        self.assertNotEqual(first, second)
        self.assertTrue(check_password("Güvenli-parola-123", first))
        self.assertFalse(check_password("yanlış", first))
        self.assertFalse(check_password("test", "bozuk"))

    def test_password_length(self):
        with self.assertRaises(ValueError):
            hash_password("123")

    def test_session_hash(self):
        self.assertEqual(len(token_hash("opaque-session")), 64)
        self.assertNotEqual(token_hash("a"), token_hash("b"))

    def test_valid_citation(self):
        self.assertTrue(valid_citations("Kaynaklı cevap. [K1]", ["K1"], ["K1", "K2"]))

    def test_unknown_citation_rejected(self):
        self.assertFalse(valid_citations("Cevap. [K99]", ["K99"], ["K1"]))

    def test_missing_inline_citation_rejected(self):
        self.assertFalse(valid_citations("Cevap.", ["K1"], ["K1"]))

    def test_metadata_citation_mismatch(self):
        self.assertFalse(valid_citations("Cevap. [K1]", ["K2"], ["K1", "K2"]))

    def test_common_local_model_citation_shapes_are_normalized(self):
        self.assertEqual(
            _normalize_citation_shapes("Cevap. (K1) [K2, K3] 【K4】"),
            "Cevap. [K1] [K2] [K3] [K4]",
        )
        self.assertEqual(
            _declared_source_ids(["[K1, K2]", "K3 ve K4"]),
            ({"K1", "K2", "K3", "K4"}, False),
        )


if __name__ == "__main__":
    unittest.main()
